"""Regenerate docs/talk-slides.md from the deck's slide text and speaker notes.

The PowerPoint deck is the source of truth; run this after editing it:
    ~/.cache/deckvenv/bin/python talk/deck_to_markdown.py [path/to/deck.pptx]
"""
import pathlib
import sys

from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parent.parent
DECK = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else
                    "~/Documents/Talks/github-universe-26/U26_whos_watching_agent_v2.pptx").expanduser()
SKIP = {"@doniacld", "2026 GitHub presentation template instructions 1"}

CHEAT_SHEET = """
## Appendix: demo cheat sheet (not a slide)

Keep this open on a second screen. Full detail in [`demo-runbook.md`](demo-runbook.md).

| Beat | Command | Point at |
|------|---------|----------|
| Setup (T-30 min) | `make demo-up` + `make forward` + `make open` | site loads, 3 pods Running |
| Panes | `make watch` · `make logs SVC=payments` · command pane | start them before uploading |
| 0 · healthy | browser: buy a ticket ($10, prefilled card) | `charged $10.00 … POST /charge`, `make watch` quiet |
| 1 · monitor | browser: upload a sighting | photo on the wall, `make watch` connect line, payments `GET /vault` |
| 2 · block | `make block` → upload → `make prove` → `make status` | 503, no `GET /vault`, policy that fired, NENFORCE 0 → 1 |
| 3 · kill | `make kill` → upload ×3 (or `make curl`) | uploads 503, gallery RESTARTS climb, frontend/payments Running |
| Fix | `git diff after-party fix/clean-upload -- services/gallery/main.go` | one import line |
| Reset | `make monitor` + `kubectl --context kind-agents-demo -n party rollout restart deploy/gallery deploy/payments` | clean wall, ticket counter back to OCTO-0001 |

Rehearsal reminders:
- `tetra getevents` is a live stream: start `make watch` *before* uploading (`make events` dumps earlier ones).
- Give `make block` / `make kill` a couple of seconds to propagate before uploading.
- `KPROBE_ACTION_OVERRIDE` in `make prove` is the policy's configured action (it shows in monitor mode too):
  the proof of the block is the 503 + `make status` NENFORCE.
- `Override` only loads on syscalls / `security_` hooks, hence `security_socket_connect` for the block
  beat and `tcp_connect` for the kill beat.
"""


def main():
    prs = Presentation(DECK)
    out = ["# Who's watching the agents? — slide deck", "",
           "Generated from the PowerPoint deck by `talk/deck_to_markdown.py`: the deck is the source",
           "of truth, edit it and re-run the script. Each slide lists its timing (start · duration),",
           "what's on screen, and the script from the speaker notes.", ""]
    for i, s in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text.replace("\x0b", " ").strip() for sh in s.shapes if sh.has_text_frame]
        texts = [t for t in texts if t and t not in SKIP and not t.startswith("TODO")]
        todos = [sh.text_frame.text.strip() for sh in s.shapes
                 if sh.has_text_frame and sh.text_frame.text.startswith("TODO")]
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        heading = texts[0].split("\n")[0] if texts else "(visual only)"
        timing, _, script = notes.partition("\n")
        out += [f"## {i} · {heading}", "", f"**Timing:** {timing or '—'}", ""]
        if len(texts) > 1 or "\n" in (texts[0] if texts else ""):
            on_screen = " · ".join(t.split("\n")[0] for t in texts[1:]) if len(texts) > 1 else ""
            if on_screen:
                out += [f"**On screen:** {on_screen}", ""]
        out += [f"**{t}**" for t in todos] + ([""] if todos else [])
        if script.strip():
            out += [line if line.startswith(("1.", "2.", "3.", "4.", "5.", "Screen:", "After", "Careful", "The repo"))
                    else f"> {line}" for line in script.strip().splitlines()] + [""]
        out += ["---", ""]
    (ROOT / "docs" / "talk-slides.md").write_text("\n".join(out) + CHEAT_SHEET)
    print("wrote docs/talk-slides.md from", DECK.name)


main()
